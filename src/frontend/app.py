"""Streamlit interface for the Finnie educational finance assistant."""

from __future__ import annotations

import base64
import logging
import sys
from io import BytesIO
from pathlib import Path
from typing import Any

import streamlit as st

# Streamlit may launch this file with src/frontend as the script directory.
# Add the repository root so imports such as `src.core.workflow` resolve
# consistently when started with `streamlit run src/frontend/app.py`.
PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.core.workflow import chain

LOGGER = logging.getLogger(__name__)

APP_TITLE = "Finnie | Personal Finance Assistant"
EDUCATION_NOTICE = (
    "**Education only:** Finnie's responses are AI-generated and may be inaccurate. "
    "They are not investment, financial, tax, or legal advice. Do not use them to "
    "make investment decisions."
)

TAB_CONFIG: list[dict[str, str]] = [
    {
        "label": "Portfolio Analyze",
        "key": "portfolio",
        "intent": "portfolio",
        "description": (
            "Discuss holdings you provide. Include approximate weights or values "
            "for a useful descriptive summary. Do not include account numbers or credentials."
        ),
        "placeholder": "Describe your holdings and approximate allocation...",
        "examples": "Example: Summarize my portfolio: 50% broad-market fund, 30% bonds, 20% cash.",
    },
    {
        "label": "Market Research",
        "key": "market",
        "intent": "market",
        "description": (
            "Ask for the latest available quote for one stock ticker. Finnie's current "
            "market workflow does not support index quotes or historical price series."
        ),
        "placeholder": "Ask about one stock ticker, such as TCS...",
        "examples": "Example: What is the latest available price for TCS?",
    },
    {
        "label": "Goal Planning",
        "key": "goals",
        "intent": "goals",
        "description": (
            "Talk through a savings or budgeting goal. Share only details you are "
            "comfortable providing, such as a target, time horizon, and current savings."
        ),
        "placeholder": "Describe a savings or budgeting goal...",
        "examples": "Example: Help me plan to save INR 100,000 in two years; I have INR 20,000 saved.",
    },
    {
        "label": "General Discussion",
        "key": "general",
        "intent": "finance_qa",
        "description": (
            "Ask general personal-finance and investing education questions grounded "
            "in Finnie's knowledge base."
        ),
        "placeholder": "Ask a general finance question...",
        "examples": "Example: What is diversification?",
    },
]


def _history_key(tab_key: str) -> str:
    return f"finnie_messages_{tab_key}"


def _prepare_portfolio_attachment(uploaded_file: Any) -> dict[str, str]:
    """Validate an upload and convert it to bounded text or inline image data."""
    if uploaded_file is None:
        return {}

    file_bytes = uploaded_file.getvalue()
    if not file_bytes:
        raise ValueError("The selected file is empty.")
    if len(file_bytes) > 5 * 1024 * 1024:
        raise ValueError("Files must be 5 MB or smaller.")

    suffix = Path(uploaded_file.name).suffix.casefold()
    attachment: dict[str, str] = {"portfolio_attachment_name": uploaded_file.name}

    if suffix == ".xlsx":
        if not file_bytes.startswith(b"PK\x03\x04"):
            raise ValueError("That file does not appear to be a valid .xlsx workbook.")

        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError(
                "Excel support is not installed yet. Run `python -m pip install -r requirements.txt`."
            ) from exc

        try:
            workbook = load_workbook(
                BytesIO(file_bytes), read_only=True, data_only=True, keep_links=False
            )
        except Exception as exc:
            raise ValueError("Finnie could not read that Excel workbook.") from exc

        lines: list[str] = []
        try:
            for sheet in workbook.worksheets[:8]:
                lines.append(f"Sheet: {sheet.title}")
                rows_read = 0
                for row_number, row in enumerate(
                    sheet.iter_rows(max_col=30, values_only=True), start=1
                ):
                    if row_number > 200:
                        lines.append("[Sheet truncated after 200 rows]")
                        break

                    values = []
                    for value in row:
                        if value is None:
                            values.append("")
                        elif hasattr(value, "isoformat"):
                            values.append(str(value.isoformat()))
                        else:
                            values.append(str(value).replace("\r", " ").replace("\n", " ")[:250])

                    if any(value.strip() for value in values):
                        lines.append(f"Row {row_number}: " + " | ".join(values))
                        rows_read += 1

                    if sum(len(line) for line in lines) > 18_000:
                        lines.append("[Workbook text truncated to fit the analysis limit]")
                        break

                if rows_read == 0:
                    lines.append("[No readable cell values found on this sheet]")
                if sum(len(line) for line in lines) > 18_000:
                    break
        finally:
            workbook.close()

        extracted_text = "\n".join(lines).strip()
        if not extracted_text:
            raise ValueError("No readable cell values were found in that workbook.")
        attachment["portfolio_attachment_text"] = extracted_text[:20_000]
        return attachment

    if suffix in {".png", ".jpg", ".jpeg"}:
        if suffix == ".png":
            if not file_bytes.startswith(b"\x89PNG\r\n\x1a\n"):
                raise ValueError("That file does not appear to be a valid PNG image.")
            mime_type = "image/png"
        else:
            if not file_bytes.startswith(b"\xff\xd8\xff"):
                raise ValueError("That file does not appear to be a valid JPG image.")
            mime_type = "image/jpeg"

        attachment["portfolio_image_data"] = base64.b64encode(file_bytes).decode("ascii")
        attachment["portfolio_image_mime_type"] = mime_type
        return attachment

    raise ValueError("Supported uploads are .xlsx, .png, .jpg, and .jpeg files.")


def _render_sources(sources: Any) -> None:
    if not isinstance(sources, list) or not sources:
        return

    with st.expander("Sources", expanded=False):
        for source in sources:
            if not isinstance(source, dict):
                continue
            title = str(source.get("title") or "Source")
            url = source.get("url")
            distance = source.get("distance")
            if isinstance(url, str) and url.startswith(("https://", "http://")):
                st.link_button(title, url, use_container_width=True)
            else:
                st.markdown(f"**{title}**")
            if isinstance(distance, (int, float)):
                st.caption(f"Retrieval distance: {distance:.3f}")


def _render_market_details(tool_result: Any) -> None:
    if not isinstance(tool_result, dict) or tool_result.get("provider") != "Alpha Vantage":
        return

    symbol = tool_result.get("symbol", "")
    trading_day = tool_result.get("latest_trading_day", "")
    retrieved_at = tool_result.get("retrieved_at_utc", "")
    provider = tool_result.get("provider", "")
    st.caption(
        f"Data: {provider} · {symbol} · latest trading day: {trading_day} "
        f"· retrieved: {retrieved_at}"
    )


def _conversation_input(history: list[dict[str, Any]], prompt: str, tab_key: str) -> str:
    # Market symbol extraction operates on the latest request, so keep that
    # request clean. Other agents can use recent user turns as conversational context.
    if tab_key == "market":
        return prompt

    earlier_user_messages = [
        str(message.get("content", ""))
        for message in history
        if message.get("role") == "user" and message.get("content")
    ][-4:]
    if not earlier_user_messages:
        return prompt

    previous_turns = "\n".join(f"- {message}" for message in earlier_user_messages)
    return (
        "Conversation context from the user's earlier messages (use only when "
        "relevant):\n"
        f"{previous_turns}\n\nCurrent user request:\n{prompt}"
    )


def _render_tab(tab: Any, config: dict[str, str]) -> None:
    tab_key = config["key"]
    history_key = _history_key(tab_key)
    if history_key not in st.session_state:
        st.session_state[history_key] = []

    with tab:
        st.markdown(f"_{config['description']}_")
        st.caption(config["examples"])

        uploaded_portfolio_file = None
        if tab_key == "portfolio":
            uploaded_portfolio_file = st.file_uploader(
                "Upload a portfolio workbook or screenshot",
                type=["xlsx", "png", "jpg", "jpeg"],
                accept_multiple_files=False,
                max_upload_size=5,
                key="portfolio_file_upload",
                help="One .xlsx, .png, or .jpg file, up to 5 MB. Ask a question in the chat to analyze it.",
            )
            st.info(
                "Uploaded information is sent to Google Gemini for analysis. Redact names, "
                "account numbers, IDs, and contact details first. AI can misread figures in "
                "screenshots; verify extracted values."
            )
            if uploaded_portfolio_file is not None:
                st.caption(
                    f"Selected: {uploaded_portfolio_file.name}. It will be included with "
                    "each chat request while it remains selected in the uploader."
                )

        
        history: list[dict[str, Any]] = st.session_state[history_key]
        for message in history:
            role = message.get("role", "assistant")
            with st.chat_message(role):
                st.markdown(str(message.get("content", "")))
                if role == "assistant":
                    _render_sources(message.get("sources"))
                    _render_market_details(message.get("tool_result"))
                    if message.get("needs_clarification"):
                        st.info("Finnie needs a little more information to continue.")

        prompt = st.chat_input(
            config["placeholder"],
            key=f"chat_input_{tab_key}",
            max_chars=2000,
            submit_mode="disable",
        )
        if not prompt or not prompt.strip():
            if st.button("Clear this chat", key=f"clear_{tab_key}", type="secondary"):
                st.session_state[history_key] = []
                st.rerun()
            return

        prompt = prompt.strip()
        history.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)

        request: dict[str, Any] = {
            "user_input": _conversation_input(history[:-1], prompt, tab_key),
            "requested_intent": config["intent"],
        }
        attachment_error = None
        if tab_key == "portfolio":
            try:
                request.update(_prepare_portfolio_attachment(uploaded_portfolio_file))
            except ValueError as exc:
                attachment_error = str(exc)
        if tab_key == "market":
            for previous_message in reversed(history[:-1]):
                previous_result = previous_message.get("tool_result")
                if isinstance(previous_result, dict) and previous_result.get("symbol"):
                    request["entities"] = {"symbol": previous_result["symbol"]}
                    break

        with st.chat_message("assistant"):
            with st.spinner("Finnie is preparing an educational response..."):
                if attachment_error:
                    response = f"I couldn't use that file: {attachment_error}"
                    sources = []
                    tool_result = {}
                    needs_clarification = False
                else:
                    try:
                        result = chain.invoke(request)
                        response = str(
                            result.get("response")
                            or "Finnie could not produce a response. Please try again."
                        )
                        sources = result.get("sources", [])
                        tool_result = result.get("tool_result", {})
                        needs_clarification = bool(result.get("needs_clarification", False))
                        error = result.get("error")
                        if error:
                            st.warning(str(error))
                    except Exception:
                        LOGGER.exception("Finnie workflow failed in the %s tab", tab_key)
                        response = (
                            "Finnie couldn't complete that request right now. Please try "
                            "again later or check that the required API configuration is available."
                        )
                        sources = []
                        tool_result = {}
                        needs_clarification = False

            st.markdown(response)
            _render_sources(sources)
            _render_market_details(tool_result)
            if needs_clarification:
                st.info("Finnie needs a little more information to continue.")

        history.append(
            {
                "role": "assistant",
                "content": response,
                "sources": sources,
                "tool_result": tool_result,
                "needs_clarification": needs_clarification,
            }
        )
        # Re-render the saved conversation before placing chat_input so the
        # composer remains below the latest user and assistant messages.
        st.rerun()

        if st.button("Clear this chat", key=f"clear_{tab_key}", type="secondary"):
            st.session_state[history_key] = []
            st.rerun()


def main() -> None:
    st.set_page_config(page_title=APP_TITLE, page_icon="💰", layout="wide")

    st.title("💰 Finnie")
    st.subheader("Your AI personal finance learning assistant")
    st.write(
        "Explore general finance concepts, describe a portfolio, look up a supported "
        "stock quote, or think through a savings goal in a conversational interface."
    )
    st.warning(EDUCATION_NOTICE, icon="⚠️")
    st.caption(
        "Chats are kept in this browser session only and are cleared when the session ends. "
        "Avoid sharing account numbers, passwords, or other sensitive personal information."
    )

    tabs = st.tabs([config["label"] for config in TAB_CONFIG])
    for tab, config in zip(tabs, TAB_CONFIG):
        _render_tab(tab, config)


if __name__ == "__main__":
    main()
